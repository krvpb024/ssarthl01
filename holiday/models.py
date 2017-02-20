from django.db import models
from userprofile.models import UserProfile
from django.db.models.signals import pre_save, post_save, post_delete, pre_delete
from django.core.urlresolvers import reverse
from django.shortcuts import get_object_or_404, render, HttpResponseRedirect
from docx import Document
from ssarthl01.settings import BASE_DIR, MEDIA_ROOT
import os, calendar, json



# Create your models here.

MONTH_CHOICES = (
	('1', '1'),
	('2', '2'),
	('3', '3'),
	('4', '4'),
	('5', '5'),
	('6', '6'),
	('7', '7'),
	('8', '8'),
	('9', '9'),
	('10', '10'),
	('11', '11'),
	('12', '12'),
)

def upload_location(instance, filename):
	ext = filename.split('.')[-1]
	return "{}/{}.{}".format('holiday',str(instance.year)+str(instance.month), ext)

class HolidayMonthFromDocx(models.Model):
	year = models.PositiveIntegerField()
	month = models.CharField(max_length=20, choices=MONTH_CHOICES)
	holiday_file = models.FileField(upload_to=upload_location)

	def __str__(self):
		return self.month

	class Meta:
		ordering = ['-pk']

	def get_absolute_url(self):
		return reverse ('holiday_detail_from_docx', kwargs={'pk':self.pk})

	def create_tablemoney_month(self):
		return '?month={}&year={}&create=True'.format(self.month, self.year)

	def remove(self):
		return '?delete=True'


def get_holiday_post_save_receiver(sender, instance, created, *args, **kwargs):
	document = Document(os.path.join(MEDIA_ROOT,'holiday',str(instance.year)+str(instance.month) + '.docx'))
	table = document.tables[0]
	numdays = int(calendar.monthrange(int(instance.year)+1911, int(instance.month))[1])
	# print(numdays)
	# print(table.cell(1, 2).text)

	# 所有人員的編號list
	all_number_list = []
	all_person = UserProfile.objects.all()
	for person in all_person:
		all_number_list.append(str(person.number))

	# 所有人員的{名字：休假}dict
	holiday_dict = {}

	x = 1
	y = 2
	z = 3
	while all_number_list != []:
		print("這是列表")
		print(all_number_list)
		if table.cell(x, y).text in all_number_list: #x=1 y=2 編號
			all_number_list.remove(table.cell(x, y).text)
			x += 1
			holiday_dict[str(table.cell(x, y).text)] = [] #x=2 y=2 名字
			while numdays != 0:
				holiday_dict[str(table.cell(x, y).text)].append(table.cell(z, y).text) #z=3 y=2 休假
				z += 1
				numdays -= 1
			print(holiday_dict)
			x = 1
			z = 3
			y += 1
			numdays = int(calendar.monthrange(int(instance.year)+1911, int(instance.month))[1])
		else:
			y += 1



	if created:
		names = UserProfile.objects.all()
		month = HolidayMonthFromDocx.objects.get(month=instance.month, year=instance.year)
		year = month.year

		for name in names:
			print(holiday_dict[name.name])
			docx_holiday, create = HolidayFromDocx.objects.get_or_create(name=name, month=month, year=year, date=json.dumps(holiday_dict[name.name], ensure_ascii=False), work_day_count=holiday_dict[name.name].count(''))

post_save.connect(get_holiday_post_save_receiver, sender=HolidayMonthFromDocx)

def docx_create_table_money_post_save_receiver(sender, instance, *args, **kwargs):
	from tablemoney.models import Month
	tablemoney_month, create = Month.objects.get_or_create(month=instance.month, year=instance.year)

	month = HolidayMonthFromDocx.objects.all().order_by('pk')
	table_number = month.count()
	while table_number >5:
		month.first().delete()
		table_number = month.count()

post_save.connect(docx_create_table_money_post_save_receiver, sender=HolidayMonthFromDocx)

def docx_table_money_post_delete_receiver(sender, instance, *args, **kwargs):
	from tablemoney.models import Month
	try:
		month = Month.objects.get(month=instance.month, year=instance.year)
		month.delete()
	except:
		pass
	try:
		if instance.holiday_file:
			if os.path.isfile(instance.holiday_file.path):
				os.remove(instance.holiday_file.path)
	except:
		pass

pre_delete.connect(docx_table_money_post_delete_receiver, sender=HolidayMonthFromDocx)


class HolidayMonth(models.Model):
	year = models.PositiveIntegerField()
	month = models.CharField(max_length=20, choices=MONTH_CHOICES)
	name = models.ManyToManyField(UserProfile, through='Holiday')
	holiday_count = models.PositiveIntegerField(default=0)


	def __str__(self):
		return self.month

	def get_absolute_url(self):
		return reverse ('holiday_detail', kwargs={'pk':self.pk})

	def get_name(self):
		names = UserProfile.objects.all()
		month = HolidayMonth.objects.get(month=self.month, year=self.year)
		year = month.year
		for name in names:
			holiday, create = Holiday.objects.get_or_create(name=name, month=month, year=year)
			if holiday.name.rank == '隊員':
				holiday.holiday_count = 50
			else:
				holiday.holiday_count = self.holiday_count
			holiday.save()

	def remove(self):
		return '%s/?delete=True' %(self.pk)

	def create_tablemoney_month(self):
		return '?month={}&year={}&create=True'.format(self.month, self.year)

	class Meta:
		ordering = ['-pk']

def get_name_post_save_receiver(sender, instance, created, *args, **kwargs):
	if created:
		names = UserProfile.objects.all()
		month = HolidayMonth.objects.get(month=instance.month, year=instance.year)
		year = month.year

		for name in names:
			holiday, create = Holiday.objects.get_or_create(name=name, month=month, year=year)
			if holiday.name.rank == '隊員':
				holiday.holiday_count = 50
			else:
				holiday.holiday_count = instance.holiday_count
			holiday.save()

post_save.connect(get_name_post_save_receiver, sender=HolidayMonth)

def create_table_money_post_save_receiver(sender, instance, *args, **kwargs):
	from tablemoney.models import Month
	tablemoney_month, create = Month.objects.get_or_create(month=instance.month, year=instance.year)


post_save.connect(create_table_money_post_save_receiver, sender=HolidayMonth)

def table_money_post_delete_receiver(sender, instance, *args, **kwargs):
	from tablemoney.models import Month
	month = Month.objects.get(month=instance.month, year=instance.year)
	month.delete()

pre_delete.connect(table_money_post_delete_receiver, sender=HolidayMonth)


class Date(models.Model):
	date = models.CharField(max_length=20)
	def __str__(self):
		return self.date


class HolidayFromDocx(models.Model):
	name = models.ForeignKey(UserProfile)
	year = models.PositiveIntegerField()
	month = models.ForeignKey(HolidayMonthFromDocx)
	date = models.CharField(max_length=1000, blank=True)
	work_day_count = models.PositiveIntegerField(default=0)
	identify = models.CharField(max_length=20, blank=True)

	def __str__(self):
		return str(self.name)

def holiday_from_Docs_get_identify_pre_save_receiver(sender, instance, *args, **kwargs):
	instance.identify = str(instance.year)+str(instance.month)

pre_save.connect(holiday_from_Docs_get_identify_pre_save_receiver, sender=HolidayFromDocx)


class Holiday(models.Model):
	name = models.ForeignKey(UserProfile)
	year = models.PositiveIntegerField()
	month = models.ForeignKey(HolidayMonth)
	date = models.ManyToManyField(Date, blank=True)
	work_day_count = models.CharField(max_length=20, blank=True)
	identify = models.CharField(max_length=20)
	holiday_count = models.PositiveIntegerField(default=0)


	def __str__(self):
		return str(self.name)


	def get_work_day_count(self):

		self.work_day_count = self.date.count()

def holiday_get_identify_pre_save_receiver(sender, instance, *args, **kwargs):
	instance.identify = str(instance.year)+str(instance.month)

pre_save.connect(holiday_get_identify_pre_save_receiver, sender=Holiday)



def update_table_money_workday_count_post_save_receiver(sender, instance, *args, **kwargs):


	from tablemoney.models import TableMoney
	try:
		tablemoney = get_object_or_404(TableMoney, name=instance.name, year=instance.year, identify=instance.identify)
		tablemoney.save()
	except:
		pass

post_save.connect(update_table_money_workday_count_post_save_receiver, sender=Holiday)
